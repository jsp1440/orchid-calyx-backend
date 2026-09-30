"""Federated ingestion for Jay Pfahl / IOSPE spreadsheet extracts.

The workbook is a partner evidence source, never the canonical taxonomy.
Records are reconciled to the existing World Plants/Hassler taxon spine by
exact canonical name or Hassler-resolved synonym. Source wording, references,
image-credit fields and IOSPE editorial markers remain attributable to IOSPE.
"""

from __future__ import annotations

import hashlib
import json
from collections.abc import Iterable, Mapping
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from runtime.knowledge_graph.canonical_taxonomy import canonical_name_of
from runtime.knowledge_graph.models import Node
from runtime.knowledge_graph.publisher import (
    DomainAdapter,
    EdgeSpec,
    NodeSpec,
    PublishResult,
    canonical_key,
    publish_domain,
)

from .yong_gee import build_taxon_index, clean_html

SOURCE_NAME = "Internet Orchid Species Photo Encyclopedia"
SOURCE_KIND = "partner_compiled_resource"
SOURCE_TABLE = "federated.iospe_workbook"
DEFAULT_SHEET = "Sheet1"

EVIDENCE_FIELDS: dict[str, str] = {
    "season": "phenology",
    "temperature": "cultivation_temperature",
    "light": "cultivation_light",
    "fragrance": "scent",
    "common name": "common_name",
    "flower size": "flower_size",
    "description": "source_description",
    "synonyms": "nomenclature",
    "references": "bibliography",
}

AUXILIARY_FIELDS = (
    "main image",
    "section type",
    "section",
    "main photo credit",
    "main photo credit link",
    "main photo credit description",
    "alt images, credits, links",
)


def stable_digest(value: Mapping[str, Any]) -> str:
    payload = json.dumps(dict(value), ensure_ascii=False, sort_keys=True, default=str)
    return hashlib.sha256(payload.encode("utf-8")).hexdigest()


def source_name_parts(raw_name: Any) -> tuple[str | None, str]:
    text = clean_html(raw_name) or ""
    marker = text[:1] if text[:1] in {"!", "~"} else None
    if marker:
        text = text[1:].strip()
    return marker, canonical_name_of(text)


@dataclass(frozen=True)
class IOSPERecord:
    source_record_id: str
    row_number: int
    scientific_name: str
    editorial_marker: str | None
    source_digest: str
    raw: dict[str, Any]
    cleaned: dict[str, str | None]


@dataclass(frozen=True)
class IOSPETaxonResolution:
    source_record_id: str
    scientific_name: str
    state: str
    taxon_pk: str | None = None
    matched_label: str | None = None
    candidate_taxon_pks: tuple[str, ...] = ()


@dataclass
class IOSPEReconciliationReport:
    total_records: int = 0
    matched: int = 0
    ambiguous: int = 0
    unresolved: int = 0
    editorial_marked: int = 0
    evidence_rows: int = 0
    resolutions: list[IOSPETaxonResolution] = field(default_factory=list)

    def to_dict(self) -> dict[str, Any]:
        return {
            "source": SOURCE_NAME,
            "total_records": self.total_records,
            "matched": self.matched,
            "ambiguous": self.ambiguous,
            "unresolved": self.unresolved,
            "editorial_marked": self.editorial_marked,
            "evidence_rows": self.evidence_rows,
            "resolutions": [r.__dict__ for r in self.resolutions],
        }


def read_iospe_workbook(
    path: str | Path, sheet_name: str = DEFAULT_SHEET
) -> list[IOSPERecord]:
    workbook = load_workbook(filename=Path(path), read_only=True, data_only=True)
    if sheet_name not in workbook.sheetnames:
        raise ValueError(
            f"sheet {sheet_name!r} not found; available={workbook.sheetnames!r}"
        )
    sheet = workbook[sheet_name]
    rows = sheet.iter_rows(values_only=True)
    try:
        headers = [str(value).strip() if value is not None else "" for value in next(rows)]
    except StopIteration:
        return []
    if "name" not in headers:
        raise ValueError("IOSPE workbook is missing required name column")

    out: list[IOSPERecord] = []
    for row_number, values in enumerate(rows, start=2):
        raw = {header: value for header, value in zip(headers, values) if header}
        if not any(value not in (None, "") for value in raw.values()):
            continue
        marker, scientific_name = source_name_parts(raw.get("name"))
        if not scientific_name:
            continue

        cleaned = {field: clean_html(raw.get(field)) for field in EVIDENCE_FIELDS}
        for field in AUXILIARY_FIELDS:
            cleaned[field] = clean_html(raw.get(field))

        out.append(
            IOSPERecord(
                source_record_id=f"iospe-row:{row_number}",
                row_number=row_number,
                scientific_name=scientific_name,
                editorial_marker=marker,
                source_digest=stable_digest(raw),
                raw=raw,
                cleaned=cleaned,
            )
        )
    return out


def reconcile_iospe_records(
    records: Iterable[IOSPERecord],
    taxonomy_nodes: Iterable[Node],
) -> tuple[list[tuple[IOSPERecord, IOSPETaxonResolution]], IOSPEReconciliationReport]:
    index = build_taxon_index(taxonomy_nodes)
    pairs: list[tuple[IOSPERecord, IOSPETaxonResolution]] = []
    report = IOSPEReconciliationReport()

    for record in records:
        report.total_records += 1
        if record.editorial_marker:
            report.editorial_marked += 1
        key = canonical_name_of(record.scientific_name).casefold()
        matches = index.get(key, []) if key else []
        if len(matches) == 1:
            node = matches[0]
            resolution = IOSPETaxonResolution(
                source_record_id=record.source_record_id,
                scientific_name=record.scientific_name,
                state="matched",
                taxon_pk=str(node.source_pk),
                matched_label=node.display_label,
            )
            report.matched += 1
        elif len(matches) > 1:
            resolution = IOSPETaxonResolution(
                source_record_id=record.source_record_id,
                scientific_name=record.scientific_name,
                state="ambiguous",
                candidate_taxon_pks=tuple(sorted(str(n.source_pk) for n in matches)),
            )
            report.ambiguous += 1
        else:
            resolution = IOSPETaxonResolution(
                source_record_id=record.source_record_id,
                scientific_name=record.scientific_name,
                state="unresolved",
            )
            report.unresolved += 1
        report.resolutions.append(resolution)
        pairs.append((record, resolution))
    return pairs, report


def iospe_evidence_rows(
    reconciled: Iterable[tuple[IOSPERecord, IOSPETaxonResolution]],
) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    for record, resolution in reconciled:
        if resolution.state != "matched" or not resolution.taxon_pk:
            continue
        for field_name, evidence_type in EVIDENCE_FIELDS.items():
            excerpt = record.cleaned.get(field_name)
            if not excerpt:
                continue
            evidence_pk = f"iospe:{record.row_number}:{field_name}"
            rows.append(
                {
                    "source_pk": evidence_pk,
                    "taxon_pk": resolution.taxon_pk,
                    "title": f"{record.scientific_name} — {field_name}",
                    "claim_label": field_name,
                    "claim_id": evidence_pk,
                    "evidence_type": evidence_type,
                    "citation": SOURCE_NAME,
                    "excerpt": excerpt,
                    "review_state": "source_imported_taxonomy_matched",
                    "evidence_class": "compiled_partner_source",
                    "confidence_score": 1.0,
                    "confidence_label": "source_faithful",
                    "source_name": SOURCE_NAME,
                    "source_kind": SOURCE_KIND,
                    "source_record_id": record.source_record_id,
                    "source_digest": record.source_digest,
                    "editorial_marker": record.editorial_marker,
                    "data_owner": "Jay Pfahl",
                    "supplied_by": "Chris Howard",
                    "underlying_citation_status": (
                        "present_in_record"
                        if record.cleaned.get("references")
                        else "unresolved"
                    ),
                }
            )
    return rows


def _produce_iospe_evidence(
    rows: Iterable[dict[str, Any]],
) -> tuple[list[NodeSpec], list[EdgeSpec]]:
    nodes: list[NodeSpec] = []
    edges: list[EdgeSpec] = []
    for row in rows:
        source_pk = row["source_pk"]
        taxon_pk = row["taxon_pk"]
        payload = {
            key: row.get(key)
            for key in (
                "claim_id",
                "claim_label",
                "evidence_type",
                "citation",
                "excerpt",
                "review_state",
                "source_name",
                "source_kind",
                "source_record_id",
                "source_digest",
                "editorial_marker",
                "data_owner",
                "supplied_by",
                "underlying_citation_status",
            )
        }
        nodes.append(
            NodeSpec(
                node_type="evidence",
                source_pk=source_pk,
                display_label=row.get("title"),
                source_table=SOURCE_TABLE,
                evidence_class=row.get("evidence_class"),
                confidence_score=row.get("confidence_score"),
                confidence_label=row.get("confidence_label"),
                payload=payload,
            )
        )
        edges.append(
            EdgeSpec(
                edge_type="supported_by_evidence",
                from_key=canonical_key("taxon", taxon_pk),
                to_key=canonical_key("evidence", source_pk),
                source_table=SOURCE_TABLE,
                source_pk=source_pk,
                evidence_class=row.get("evidence_class"),
                confidence_score=row.get("confidence_score"),
                confidence_label=row.get("confidence_label"),
                rule_name="iospe_federated_ingestion",
                payload={
                    "source_name": SOURCE_NAME,
                    "source_record_id": row.get("source_record_id"),
                    "claim_label": row.get("claim_label"),
                },
            )
        )
    return nodes, edges


IOSPE_EVIDENCE_ADAPTER = DomainAdapter(
    domain="evidence",
    source_table=SOURCE_TABLE,
    produce=_produce_iospe_evidence,
    required_identifiers=("source_pk", "taxon_pk"),
)


def publish_iospe_evidence(repo: Any, rows: Iterable[dict[str, Any]]) -> PublishResult:
    return publish_domain(repo, IOSPE_EVIDENCE_ADAPTER, rows)


def build_iospe_dry_run(
    workbook_path: str | Path,
    taxonomy_nodes: Iterable[Node],
    *,
    sheet_name: str = DEFAULT_SHEET,
) -> tuple[list[dict[str, Any]], IOSPEReconciliationReport]:
    records = read_iospe_workbook(workbook_path, sheet_name=sheet_name)
    reconciled, report = reconcile_iospe_records(records, taxonomy_nodes)
    rows = iospe_evidence_rows(reconciled)
    report.evidence_rows = len(rows)
    return rows, report
