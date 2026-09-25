"""Federated ingestion for the Gary Yong Gee orchid database extract.

The workbook is a compiled specialist resource, not a canonical taxonomy.
Records are reconciled to the existing Orchid Continuum taxon spine and then
projected as provenance-preserving evidence nodes.

Nothing here mutates canonical taxonomy. Unmatched or ambiguous records remain
review items. Production publication must still pass through the existing
controlled Knowledge Graph publication workflow.
"""

from __future__ import annotations

import hashlib
import html
import json
import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Iterable, Mapping

from openpyxl import load_workbook

from runtime.knowledge_graph.canonical_taxonomy import canonical_name_of
from runtime.knowledge_graph.models import Node
from runtime.knowledge_graph.publisher import (
    DomainAdapter,
    EdgeSpec,
    NodeSpec,
    PublishResult,
    canonical_key,
    publish_domain,
)

SOURCE_NAME = "Gary Yong Gee Orchid Database"
SOURCE_KIND = "specialist_compiled_resource"
SOURCE_TABLE = "federated.gary_yong_gee_workbook"
DEFAULT_SHEET = "Chosen"
NULL_STRINGS = {"", "null", "none", "n/a", "na"}

EVIDENCE_FIELDS: dict[str, str] = {
    "publicationsp": "nomenclatural_publication",
    "pubyrsp": "publication_year",
    "etymology": "etymology",
    "synonym": "nomenclature",
    "commonName": "common_name",
    "section": "taxonomy_section",
    "subsection": "taxonomy_subsection",
    "distributionsp": "distribution",
    "habitat": "habitat",
    "season": "phenology",
    "scent": "scent",
    "characteristicsp": "morphology",
    "capsule": "fruit_capsule",
    "similarSpecies": "diagnostic_comparison",
    "notes": "taxon_notes",
    "referencesp": "bibliography",
    "notesGary": "compiler_note",
}

_BR_RE = re.compile(r"<\s*br\s*/?\s*>", re.IGNORECASE)
_BLOCK_END_RE = re.compile(r"</\s*(?:p|li|div|ul|ol)\s*>", re.IGNORECASE)
_TAG_RE = re.compile(r"<[^>]+>")
_WS_RE = re.compile(r"[ \t\f\v]+")
_MULTI_NL_RE = re.compile(r"\n{3,}")


def _none_if_null(value: Any) -> Any:
    if value is None:
        return None
    if isinstance(value, str) and value.strip().casefold() in NULL_STRINGS:
        return None
    return value


def clean_html(value: Any) -> str | None:
    value = _none_if_null(value)
    if value is None:
        return None
    text = str(value)
    text = _BR_RE.sub("\n", text)
    text = _BLOCK_END_RE.sub("\n", text)
    text = _TAG_RE.sub("", text)
    text = html.unescape(text).replace("\xa0", " ")
    lines = [_WS_RE.sub(" ", line).strip() for line in text.splitlines()]
    text = "\n".join(line for line in lines if line)
    return _MULTI_NL_RE.sub("\n\n", text).strip() or None


def stable_digest(value: Mapping[str, Any]) -> str:
    payload = json.dumps(dict(value), ensure_ascii=False, sort_keys=True, default=str)
    return hashlib.sha256(payload.encode("utf-8")).hexdigest()


def scientific_name_from_row(row: Mapping[str, Any]) -> str:
    informal = clean_html(row.get("websiteInformalName"))
    if informal:
        return canonical_name_of(informal)

    formal = clean_html(row.get("websiteFormalName"))
    if formal:
        return canonical_name_of(formal)

    genus = clean_html(row.get("genus"))
    species = clean_html(row.get("speciesName"))
    if genus and species:
        parts = [genus, species]
        for rank_field, epithet_field in (("subtax1", "taxon1"), ("subtax2", "taxon2")):
            rank = clean_html(row.get(rank_field))
            epithet = clean_html(row.get(epithet_field))
            if rank and epithet:
                parts.extend([rank, epithet])
        return canonical_name_of(" ".join(parts))
    return ""


@dataclass(frozen=True)
class YongGeeRecord:
    source_record_id: str
    scientific_name: str
    source_digest: str
    raw: dict[str, Any]
    cleaned: dict[str, str | None]


@dataclass(frozen=True)
class TaxonResolution:
    source_record_id: str
    scientific_name: str
    state: str
    taxon_pk: str | None = None
    matched_label: str | None = None
    candidate_taxon_pks: tuple[str, ...] = ()


@dataclass
class ReconciliationReport:
    total_records: int = 0
    matched: int = 0
    ambiguous: int = 0
    unresolved: int = 0
    evidence_rows: int = 0
    resolutions: list[TaxonResolution] = field(default_factory=list)

    def to_dict(self) -> dict[str, Any]:
        return {
            "source": SOURCE_NAME,
            "total_records": self.total_records,
            "matched": self.matched,
            "ambiguous": self.ambiguous,
            "unresolved": self.unresolved,
            "evidence_rows": self.evidence_rows,
            "resolutions": [r.__dict__ for r in self.resolutions],
        }


def read_workbook(path: str | Path, sheet_name: str = DEFAULT_SHEET) -> list[YongGeeRecord]:
    workbook = load_workbook(filename=Path(path), read_only=True, data_only=True)
    if sheet_name not in workbook.sheetnames:
        raise ValueError(
            f"sheet {sheet_name!r} not found; available={workbook.sheetnames!r}"
        )
    sheet = workbook[sheet_name]
    rows = sheet.iter_rows(values_only=True)
    try:
        headers = [str(value).strip() if value is not None else "" for value in next(rows)]
    except StopIteration:
        return []
    if "id" not in headers or "speciesName" not in headers:
        raise ValueError("Yong Gee workbook is missing required id/speciesName columns")

    out: list[YongGeeRecord] = []
    for values in rows:
        raw = {
            header: _none_if_null(value)
            for header, value in zip(headers, values)
            if header
        }
        source_id = str(raw.get("id") or "").strip()
        if not source_id:
            continue
        scientific_name = scientific_name_from_row(raw)
        cleaned = {field: clean_html(raw.get(field)) for field in EVIDENCE_FIELDS}
        for field in (
            "author",
            "author1",
            "author2",
            "websiteInformalName",
            "websiteFormalName",
            "websiteSlug",
            "websiteIsPublished",
            "websiteUpdatedAt",
        ):
            cleaned[field] = clean_html(raw.get(field))
        out.append(
            YongGeeRecord(
                source_record_id=source_id,
                scientific_name=scientific_name,
                source_digest=stable_digest(raw),
                raw=raw,
                cleaned=cleaned,
            )
        )
    return out


def _taxon_names(node: Node) -> set[str]:
    names: set[str] = set()
    for value in (
        node.display_label,
        node.payload.get("scientific_name"),
        node.payload.get("canonical_name"),
        node.payload.get("accepted_name"),
    ):
        if value:
            name = canonical_name_of(str(value))
            if name:
                names.add(name.casefold())
    for value in node.payload.get("synonyms", []) or []:
        name = canonical_name_of(str(value))
        if name:
            names.add(name.casefold())
    return names


def build_taxon_index(taxonomy_nodes: Iterable[Node]) -> dict[str, list[Node]]:
    index: dict[str, list[Node]] = {}
    for node in taxonomy_nodes:
        if node.node_type != "taxon":
            continue
        for name in _taxon_names(node):
            index.setdefault(name, []).append(node)
    return index


def reconcile_records(
    records: Iterable[YongGeeRecord],
    taxonomy_nodes: Iterable[Node],
) -> tuple[list[tuple[YongGeeRecord, TaxonResolution]], ReconciliationReport]:
    """Resolve by exact canonical name/synonym only; fuzzy matches require review."""

    index = build_taxon_index(taxonomy_nodes)
    pairs: list[tuple[YongGeeRecord, TaxonResolution]] = []
    report = ReconciliationReport()
    for record in records:
        report.total_records += 1
        key = canonical_name_of(record.scientific_name).casefold()
        matches = index.get(key, []) if key else []
        if len(matches) == 1:
            node = matches[0]
            resolution = TaxonResolution(
                source_record_id=record.source_record_id,
                scientific_name=record.scientific_name,
                state="matched",
                taxon_pk=str(node.source_pk),
                matched_label=node.display_label,
            )
            report.matched += 1
        elif len(matches) > 1:
            resolution = TaxonResolution(
                source_record_id=record.source_record_id,
                scientific_name=record.scientific_name,
                state="ambiguous",
                candidate_taxon_pks=tuple(sorted(str(n.source_pk) for n in matches)),
            )
            report.ambiguous += 1
        else:
            resolution = TaxonResolution(
                source_record_id=record.source_record_id,
                scientific_name=record.scientific_name,
                state="unresolved",
            )
            report.unresolved += 1
        report.resolutions.append(resolution)
        pairs.append((record, resolution))
    return pairs, report


def evidence_rows(
    reconciled: Iterable[tuple[YongGeeRecord, TaxonResolution]],
) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    for record, resolution in reconciled:
        if resolution.state != "matched" or not resolution.taxon_pk:
            continue
        slug = record.cleaned.get("websiteSlug")
        source_uri = f"https://www.yonggee.name/{slug}" if slug else None
        for field_name, evidence_type in EVIDENCE_FIELDS.items():
            excerpt = record.cleaned.get(field_name)
            if not excerpt:
                continue
            evidence_pk = f"yong-gee:{record.source_record_id}:{field_name}"
            rows.append(
                {
                    "source_pk": evidence_pk,
                    "taxon_pk": resolution.taxon_pk,
                    "title": f"{record.scientific_name} — {field_name}",
                    "claim_label": field_name,
                    "claim_id": evidence_pk,
                    "evidence_type": evidence_type,
                    "citation": SOURCE_NAME,
                    "source_uri": source_uri,
                    "excerpt": excerpt,
                    "review_state": "source_imported_taxonomy_matched",
                    "evidence_class": "compiled_specialist_source",
                    "confidence_score": 1.0,
                    "confidence_label": "source_faithful",
                    "source_name": SOURCE_NAME,
                    "source_kind": SOURCE_KIND,
                    "source_record_id": record.source_record_id,
                    "source_digest": record.source_digest,
                    "compiler": "Gary Yong Gee",
                    "underlying_citation_status": (
                        "present_in_record"
                        if record.cleaned.get("referencesp")
                        else "unresolved"
                    ),
                }
            )
    return rows


def _produce_yong_gee_evidence(
    rows: Iterable[dict[str, Any]],
) -> tuple[list[NodeSpec], list[EdgeSpec]]:
    nodes: list[NodeSpec] = []
    edges: list[EdgeSpec] = []
    for row in rows:
        source_pk = row["source_pk"]
        taxon_pk = row["taxon_pk"]
        payload = {
            key: row.get(key)
            for key in (
                "claim_id",
                "claim_label",
                "evidence_type",
                "citation",
                "source_uri",
                "excerpt",
                "review_state",
                "source_name",
                "source_kind",
                "source_record_id",
                "source_digest",
                "compiler",
                "underlying_citation_status",
            )
        }
        nodes.append(
            NodeSpec(
                node_type="evidence",
                source_pk=source_pk,
                display_label=row.get("title"),
                source_table=SOURCE_TABLE,
                evidence_class=row.get("evidence_class"),
                confidence_score=row.get("confidence_score"),
                confidence_label=row.get("confidence_label"),
                payload=payload,
            )
        )
        edges.append(
            EdgeSpec(
                edge_type="supported_by_evidence",
                from_key=canonical_key("taxon", taxon_pk),
                to_key=canonical_key("evidence", source_pk),
                source_table=SOURCE_TABLE,
                source_pk=source_pk,
                evidence_class=row.get("evidence_class"),
                confidence_score=row.get("confidence_score"),
                confidence_label=row.get("confidence_label"),
                rule_name="yong_gee_federated_ingestion",
                payload={
                    "source_name": SOURCE_NAME,
                    "source_record_id": row.get("source_record_id"),
                    "claim_label": row.get("claim_label"),
                },
            )
        )
    return nodes, edges


YONG_GEE_EVIDENCE_ADAPTER = DomainAdapter(
    domain="evidence",
    source_table=SOURCE_TABLE,
    produce=_produce_yong_gee_evidence,
    required_identifiers=("source_pk", "taxon_pk"),
)


def publish_matched_evidence(repo: Any, rows: Iterable[dict[str, Any]]) -> PublishResult:
    """Attach matched evidence through the existing KG publisher."""

    return publish_domain(repo, YONG_GEE_EVIDENCE_ADAPTER, rows)


def build_dry_run(
    workbook_path: str | Path,
    taxonomy_nodes: Iterable[Node],
    *,
    sheet_name: str = DEFAULT_SHEET,
) -> tuple[list[dict[str, Any]], ReconciliationReport]:
    records = read_workbook(workbook_path, sheet_name=sheet_name)
    reconciled, report = reconcile_records(records, taxonomy_nodes)
    rows = evidence_rows(reconciled)
    report.evidence_rows = len(rows)
    return rows, report
